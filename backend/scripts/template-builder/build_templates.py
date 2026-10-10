"""
Membangun 4 template laporan ber-placeholder (docxtemplater) dari dokumen asli klinik.
Kop surat, header/footer, warna sel, dan tata letak dipertahankan; data klien contoh dihapus.

Pemakaian: python3 build_templates.py <folder_sumber_docx> <folder_output>
"""
import copy, re, sys, os, glob
import docx
from docx.oxml.ns import qn

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

def runs(p): return p.findall(qn('w:r'))
def ptext(p): return ''.join(t.text or '' for t in p.iter(qn('w:t')))

def make_run(text, rpr=None):
    r = docx.oxml.OxmlElement('w:r')
    if rpr is not None: r.append(copy.deepcopy(rpr))
    t = docx.oxml.OxmlElement('w:t'); t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve'); t.text = text
    r.append(t); return r

def run_rpr(r): return r.find(qn('w:rPr'))

def set_par_text(p, text):
    """Ganti seluruh teks paragraf, format run pertama dipertahankan."""
    rs = runs(p)
    rpr = run_rpr(rs[0]) if rs else (p.find(qn('w:pPr')).find(qn('w:rPr')) if p.find(qn('w:pPr')) is not None else None)
    for r in rs: p.remove(r)
    for h in p.findall(qn('w:hyperlink')): p.remove(h)
    p.append(make_run(text, rpr))

def append_par(p, text):
    rs = runs(p)
    p.append(make_run(text, run_rpr(rs[-1]) if rs else None))

def fill_after_colon(p, placeholder):
    """'Label\t:  \tNILAI' -> 'Label\t:  \t{placeholder}'"""
    rs = runs(p); acc = ''
    for i, r in enumerate(rs):
        t_el = r.find(qn('w:t'))
        txt = t_el.text if t_el is not None and t_el.text else ''
        if ':' in txt:
            idx = txt.index(':')
            rest = txt[idx + 1:]
            keep_ws = re.match(r'^[\s\t]*', rest).group(0)
            t_el.text = txt[:idx + 1] + (keep_ws if keep_ws else ' ')
            t_el.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
            tail_rpr = run_rpr(r)
            for r2 in rs[i + 1:]:
                if r2.find(qn('w:tab')) is not None and not ptext_run(r2).strip():
                    continue  # pertahankan tab pemisah
                p.remove(r2)
            p.append(make_run(placeholder, tail_rpr)); return
    append_par(p, ' ' + placeholder)

def ptext_run(r): return ''.join(t.text or '' for t in r.iter(qn('w:t')))

def cell_pars(tc): return tc.findall(qn('w:p'))

def set_cell(tc, text, keep_paragraphs=1):
    ps = cell_pars(tc)
    set_par_text(ps[0], text)
    for extra in ps[keep_paragraphs:]: tc.remove(extra)

def rows(tbl): return tbl.findall(qn('w:tr'))
def cells(tr): return tr.findall(qn('w:tc'))
def remove_rows(tbl, idxs):
    rs = rows(tbl)
    for i in sorted(idxs, reverse=True): tbl.remove(rs[i])

def body_children(d): return [c for c in d.element.body.iterchildren() if not c.tag.endswith('sectPr')]

def scale_cells(tr, start, n, close=''):
    cs = cells(tr)[start:start + n]
    for k, tc in enumerate(cs, 1):
        txt = '{c%d}' % k
        if k == n: txt += close
        set_cell(tc, txt)

def strip_footer_header_names(d):
    pass  # header/footer hanya berisi kop & nomor halaman (sudah dicek)

# ------------------------------------------------------------------ TES ANAK
def tes_anak(src, out):
    d = docx.Document(src); ch = body_children(d)
    ident = {4: '{nama}', 5: '{tanggal_lahir}', 6: '{tanggal_tes}', 7: '{usia}', 8: '{alamat}', 9: '{sekolah}',
             10: '{ayah_nama}', 11: '{ayah_pendidikan}', 12: '{ayah_pekerjaan}', 13: '{ibu_nama}',
             14: '{ibu_pendidikan}', 15: '{ibu_pekerjaan}', 16: '{alamat_orang_tua}'}
    for i, ph in ident.items(): fill_after_colon(ch[i], ph)
    tbl = ch[18]; rs = rows(tbl)
    set_cell(cells(rs[5])[0], '{#kelompok}{nama_kelompok}')
    item = cells(rs[6])
    set_cell(item[0], '{#items}{no}.'); set_cell(item[1], '{label}'); set_cell(item[2], '{deskripsi}')
    scale_cells(rs[6], 3, 6, '{/items}{/kelompok}')
    # baris IQ & saran
    iq = cell_pars(cells(rs[22])[0])
    fill_after_colon(iq[1], '{iq}'); fill_after_colon(iq[2], '{skala_iq}')
    set_par_text(iq[3], 'KESIMPULAN : {kategori_iq}')
    saran_tc = cells(rs[23])[0]; sp = cell_pars(saran_tc)
    newp = copy.deepcopy(sp[0]); set_par_text(newp, '{saran}'); sp[-1].addnext(newp)
    remove_rows(tbl, [4] + list(range(7, 22)))
    t2 = ch[21]; r1 = cells(rows(t2)[1])
    set_cell(r1[0], '{kelebihan}'); set_cell(r1[1], '{kekurangan}')
    d.save(out)

# ------------------------------------------------------------------ MINAT BAKAT
def minat_bakat(src, out):
    d = docx.Document(src); ch = body_children(d)
    # halaman sampul: kotak identitas berisi data klien contoh -> placeholder
    cover = ch[11]; tc = cells(rows(cover)[0])[0]
    phs = ['{nama}', '{tanggal_lahir}', '{tanggal_tes}', '{usia}', '{pendidikan}']
    k = 0
    for p in cell_pars(tc):
        if ':' in ptext(p) and k < len(phs): fill_after_colon(p, phs[k]); k += 1
    for i, ph in zip([35, 36, 37, 38, 39], phs): fill_after_colon(ch[i], ph)
    tbl = ch[41]; rs = rows(tbl)
    # baris 2 (kecerdasan umum) memiliki sel skala tergabung -> pakai baris 3 sebagai baris loop
    it = cells(rs[3]); ps = cell_pars(it[1])
    set_cell(it[0], '{#items}{no}.')
    set_par_text(ps[0], '{label_kapital}'); set_par_text(ps[1], '{deskripsi}')
    for extra in ps[2:]: it[1].remove(extra)
    scale_cells(rs[3], 2, 5, '{/items}')
    remove_rows(tbl, [2] + list(range(4, len(rs))))
    k_tbl = ch[43]; kr = rows(k_tbl)
    set_cell(cells(kr[1])[1], '{taraf_kecerdasan}')
    set_cell(cells(kr[2])[1], '{minat_bakat}')
    set_cell(cells(kr[3])[1], '{rekomendasi_jurusan}')
    set_cell(cells(kr[4])[1], '{rekomendasi_jurusan_kuliah}')
    set_cell(cells(kr[5])[1], '{psikolog_nama}')
    # hapus paragraf kosong setelah tabel terakhir (menyebabkan halaman kosong)
    for extra in ch[44:]:
        if not ptext(extra).strip() and extra.tag.endswith('}p'): extra.getparent().remove(extra)
    d.save(out)

# ------------------------------------------------------------------ REKRUTMEN
def rekrutmen(src, out):
    d = docx.Document(src); ch = body_children(d)
    idt = rows(ch[3])
    mapping = [('{no_tes}', '{usia}'), ('{nama}', '{pendidikan}'), ('{tanggal_lahir}', '{tanggal_tes}'), ('{jenis_kelamin}', '{jabatan}')]
    for tr, (a, b) in zip(idt, mapping):
        cs = cells(tr)
        fill_after_colon(cell_pars(cs[0])[0], a); fill_after_colon(cell_pars(cs[1])[0], b)
    for idx, grp in [(8, 'grup_aspek_psikologis'), (10, 'grup_kompetensi_manajerial')]:
        tbl = ch[idx]; rs = rows(tbl); c = cells(rs[0])
        set_cell(c[0], '{#%s}{label}' % grp); set_cell(c[1], '{deskripsi}')
        scale_cells(rs[0], 2, 5, '{/%s}' % grp)
        remove_rows(tbl, range(1, len(rs)))
    set_par_text(ch[23], '{uraian}'); ch[24].getparent().remove(ch[24]); ch[25].getparent().remove(ch[25])
    set_par_text(ch[28], '{kelebihan}')
    set_par_text(ch[31], '{kekurangan}')
    set_par_text(ch[34], '{rekomendasi_uraian}')
    rek = rows(ch[36])
    for tr, slug in zip(rek, ['dapat_disarankan', 'masih_dapat_disarankan', 'kurang_dapat_disarankan', 'tidak_dapat_disarankan']):
        set_cell(cells(tr)[0], '{p_rekomendasi_%s}' % slug)
    set_par_text(ch[38], '{kota}, {tanggal_laporan}')
    set_par_text(ch[44], '{psikolog_nama}')
    set_par_text(ch[45], '                   SIPP: {psikolog_sipp}')
    d.save(out)

# ------------------------------------------------------------------ PSIKOLOGI UMUM
def psikologi_umum(src, out):
    d = docx.Document(src); ch = body_children(d)
    for i, ph in zip([2, 3, 4, 5, 6], ['{nama}', '{tanggal_lahir}', '{tanggal_tes}', '{usia}', '{pendidikan}']):
        fill_after_colon(ch[i], ph)
    tbl = ch[8]; rs = rows(tbl)
    set_cell(cells(rs[3])[2], '{general_intelligence}')
    set_cell(cells(rs[4])[0], '{#kelompok}{nama_kelompok}')
    it = cells(rs[5]); ps = cell_pars(it[1])
    set_cell(it[0], '{#items}{no}.')
    set_par_text(ps[0], '{label_kapital}'); set_par_text(ps[1], '{deskripsi}')
    scale_cells(rs[5], 2, 5, '{/items}{/kelompok}')
    remove_rows(tbl, range(6, len(rs)))
    ks = cells(rows(ch[10])[1])[0]; kp = cell_pars(ks)
    set_par_text(kp[0], '{kesimpulan}')
    for extra in kp[3:]: ks.remove(extra)
    t = rows(ch[12])
    set_cell(cells(t[0])[5], '{kota}, {tanggal_laporan}')
    for tr, slug in zip(t[1:4], ['disarankan', 'masih_disarankan', 'tidak_disarankan']):
        set_cell(cells(tr)[0], '{p_rekomendasi_%s}' % slug)
    set_par_text(ch[15], '{psikolog_nama}')
    set_par_text(ch[16], 'SIPP: {psikolog_sipp}')
    set_par_text(ch[17], '{#psikolog_str}STR Psikolog Klinis: {psikolog_str}{/psikolog_str}')
    d.save(out)

if __name__ == '__main__':
    src_dir, out_dir = sys.argv[1], sys.argv[2]
    os.makedirs(out_dir, exist_ok=True)
    find = lambda pat: glob.glob(os.path.join(src_dir, pat))[0]
    tes_anak(find('*Draft_Laporan_Tes_Anak.docx'), os.path.join(out_dir, 'TES_ANAK.docx'))
    minat_bakat(find('*Draft_Laporan_Tes_Minat_Bakat.docx'), os.path.join(out_dir, 'MINAT_BAKAT.docx'))
    rekrutmen(find('*Draft_Rekrutmen.docx'), os.path.join(out_dir, 'REKRUTMEN.docx'))
    psikologi_umum(find('*Template_Laporan_Psikologi.docx'), os.path.join(out_dir, 'PSIKOLOGI_UMUM.docx'))
    print('Template dibuat di', out_dir)
